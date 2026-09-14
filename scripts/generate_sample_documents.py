#!/usr/bin/env python3
"""Generate binary-format sample documents (PDF, DOCX, XLSX) for the demo corpus.

Run from anywhere; paths are relative to the repo root:
    python scripts/generate_sample_documents.py
"""
from __future__ import annotations

import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "data" / "documents"
DOCS.mkdir(parents=True, exist_ok=True)

SECTION_RE = re.compile(r"^(?P<num>\d+(?:\.\d+)*)[.)]?\s+(?P<title>[A-Z][^\n]{2,90})$")


def read_doc(path: Path) -> tuple[str, list[tuple[str, str]]]:
    """Very small md -> (title, [(section, body)]) reader for our own samples."""
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    title = lines[0].lstrip("# ").strip() if lines else path.name
    sections: list[tuple[str, str]] = []
    cur_title, buf = "Details", []
    for ln in lines[1:]:
        m = SECTION_RE.match(ln.strip())
        if m:
            if any(b.strip() for b in buf):
                sections.append((cur_title, "\n".join(buf).strip()))
            cur_title, buf = f"{m.group('num')} {m.group('title')}", []
        else:
            buf.append(ln)
    if any(b.strip() for b in buf):
        sections.append((cur_title, "\n".join(buf).strip()))
    return title, sections


def build_pdf(title: str, sections: list[tuple[str, str]], out: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    h = ParagraphStyle("h2x", parent=styles["Heading2"], fontSize=12, spaceAfter=6)
    b = ParagraphStyle("bodyx", parent=styles["BodyText"], fontSize=9.5, leading=13)
    doc = SimpleDocTemplate(str(out), pagesize=A4, title=title, author="Operations Team")
    story = [Paragraph(title.replace("_", " "), styles["Title"]), Spacer(1, 10)]
    for sec, body in sections:
        story.append(Paragraph(sec, h))
        for para in body.split("\n"):
            if para.strip():
                story.append(Paragraph(para.strip(), b))
        story.append(Spacer(1, 6))
    doc.build(story)


def build_docx(title: str, sections: list[tuple[str, str]], out: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading(title.replace("_", " "), level=0)
    for sec, body in sections:
        d.add_heading(sec, level=1)
        for para in body.split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
    d.core_properties.title = title
    d.core_properties.author = "Operations Team"
    d.save(out)


def main() -> None:
    pairs = [
        ("RUNBOOK-005_Payment_API_Runbook.md", "RUNBOOK-005_Payment_API_Runbook.pdf"),
        ("SOP-API-01_Deployment_Validation.txt", "SOP-API-01_Deployment_Validation.pdf"),
        ("RCA-018_Order_Service_DB_Timeout.md", "RCA-018_Order_Service_DB_Timeout.docx"),
        ("INC-044_Pod_OOMKilled_Postmortem.md", "INC-044_Pod_OOMKilled_Postmortem.docx"),
    ]
    for src, dst in pairs:
        sp = DOCS / src
        if not sp.exists():
            print(f"skip (missing source): {src}")
            continue
        title, sections = read_doc(sp)
        out = DOCS / dst
        if dst.endswith(".pdf"):
            build_pdf(title, sections, out)
        else:
            build_docx(title, sections, out)
        print(f"generated {out.name} ({out.stat().st_size} bytes)")

    # XLSX: escalation matrix -------------------------------------------------
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Escalation Matrix"
    ws.append(["Severity", "First Responder", "Escalate After (min)", "Escalation Path", "War Room Channel"])
    ws.append(["P0", "On-call SRE", 5, "SRE Lead -> Platform Director -> VP Eng", "#war-room-p0"])
    ws.append(["P1", "On-call SRE", 15, "SRE Lead -> Service Owner", "#inc-payments"])
    ws.append(["P2", "Service Engineer", 60, "Team Lead -> SRE Lead", "#inc-general"])
    ws.append(["P3", "Service Engineer", 240, "Team backlog", "-"])
    ws2 = wb.create_sheet("Service Owners")
    ws2.append(["Service", "Primary Owner", "Backup", "SLO (availability)", "Runbook Ref"])
    ws2.append(["Payment API", "payments-core@corp", "sre-oncall@corp", 0.9995, "RUNBOOK-005"])
    ws2.append(["Order Service", "orders@corp", "sre-oncall@corp", 0.999, "RCA-018"])
    ws2.append(["User Service", "identity@corp", "sre-oncall@corp", 0.999, "INC-044"])
    ws2.append(["Auth Service", "identity@corp", "secops@corp", 0.9999, "SOP-API-01"])
    out = DOCS / "SRE_Escalation_Matrix.xlsx"
    wb.save(out)
    print(f"generated {out.name} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
