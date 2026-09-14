"""Multi-format ingestion: parsers, chunking, metadata, error handling."""
from __future__ import annotations

import json

import pytest

from app.rag.ingestion.chunking import chunk_document_text, clean_text, split_sections
from app.rag.ingestion.parsers import (
    DocumentParseError, parse_csv, parse_docx, parse_json, parse_log, parse_markdown,
    parse_pdf, parse_text, parse_xlsx, parse_bytes,
)

LONG_TEXT = (
    "# Payment API Runbook\n\n"
    "## 1 Diagnosis\n\nCheck deployment status. Inspect application logs after deployment. "
    "Verify ConfigMap availability in the production namespace with kubectl.\n\n"
    "## 2 Recovery\n\nRestart affected pods. Validate Kubernetes service endpoints. "
    "Roll back the release if error rate stays above threshold.\n"
)


def test_clean_text_normalises():
    raw = "Hello\r\nWorld\x00\x01   spaced\ttabs\r\n\r\n\r\nend"
    cleaned = clean_text(raw)
    assert "World" in cleaned and "\x00" not in cleaned
    assert "   " not in cleaned


def test_sections_and_chunks_keep_provenance():
    chunks = chunk_document_text(LONG_TEXT, chunk_size=220, chunk_overlap=40)
    assert chunks, "expected at least one chunk"
    sections = {s for _, s, _ in chunks}
    assert any("1 Diagnosis" in (s or "") for s in sections)
    for idx, section, text in chunks:
        assert isinstance(idx, int) and len(text) > 20


def test_split_sections_counts():
    secs = dict(split_sections(LONG_TEXT))
    assert any("Diagnosis" in k for k in secs)


def test_markdown_and_text_parsers():
    pf = parse_markdown(LONG_TEXT.encode(), "rb.md")
    assert "ConfigMap" in pf.text
    assert parse_text(b"plain content " * 10, "a.txt").text.startswith("plain")


def test_csv_parser_produces_records_and_incident_hint():
    csv_data = "id,title,service,root_cause\nINC-1,DB timeout,Order Service,pool exhaustion\nINC-2,503,Payment API,missing config\n"
    pf = parse_csv(csv_data.encode(), "incidents.csv")
    assert "pool exhaustion" in pf.text
    recs = pf.hints["incident_records"]
    assert recs[0]["id"] == "INC-1"


def test_json_parser_recognises_incidents():
    data = json.dumps({"incidents": [{"id": "INC-9", "title": "disk full", "root_cause": "log rotation disabled"}]}).encode()
    pf = parse_json(data, "x.json")
    assert pf.hints["incident_records"][0]["id"] == "INC-9"


def test_json_invalid_raises():
    with pytest.raises(DocumentParseError):
        parse_json(b"{not json", "x.json")


def test_log_parser_filters_health_noise():
    lines = ["2026-01-01T00:00:00Z ERROR app db connection refused timeout"] * 10 + ["2026-01-01T00:00:01Z INFO GET /healthz 200"] * 50
    pf = parse_log("\n".join(lines).encode(), "app.log")
    assert "refused" in pf.text and "/healthz" not in pf.text


def test_pdf_and_docx_and_xlsx_roundtrip(tmp_path):
    from docx import Document as Docx
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas

    # PDF
    pdf_path = tmp_path / "r.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(60, 780, "Payment API Runbook")
    c.drawString(60, 760, "1 Diagnosis: Verify ConfigMap availability with kubectl get configmap")
    c.drawString(60, 740, "2 Recovery: Restart affected pods and validate endpoints after the fix")
    c.save()
    pf = parse_pdf(pdf_path.read_bytes(), "r.pdf")
    assert "ConfigMap" in pf.text and pf.page_count == 1

    # DOCX
    docx_path = tmp_path / "r.docx"
    d = Docx()
    d.add_heading("Order Service RCA", level=1)
    d.add_paragraph("Root cause: database connection pool exhaustion under peak traffic load.")
    d.save(str(docx_path))
    pf = parse_docx(docx_path.read_bytes(), "r.docx")
    assert "pool exhaustion" in pf.text

    # XLSX
    xlsx_path = tmp_path / "m.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Service", "Owner"])
    ws.append(["Payment API", "payments-team"])
    wb.save(xlsx_path)
    pf = parse_xlsx(xlsx_path.read_bytes(), "m.xlsx")
    assert "payments-team" in pf.text


def test_unsupported_and_corrupted_rejected():
    with pytest.raises(DocumentParseError):
        parse_bytes(b"MZ\x90\x00", "evil.exe")
    with pytest.raises(DocumentParseError):
        parse_pdf(b"not a real pdf at all", "broken.pdf")


def test_zip_bomb_guard():
    import io
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("inner.txt", "x" * 100)
    with pytest.raises(DocumentParseError):
        parse_bytes(buf.getvalue(), "payload.zip")


def test_upload_validation_errors(client, tokens):
    # unsupported type
    r = client.post("/api/documents/upload", headers=tokens["admin_h"], files={"file": ("a.exe", b"MZ123", "application/octet-stream")})
    assert r.status_code == 415
    # empty file
    r = client.post("/api/documents/upload", headers=tokens["admin_h"], files={"file": ("e.txt", b"", "text/plain")})
    assert r.status_code == 400
    # too little text
    r = client.post("/api/documents/upload", headers=tokens["admin_h"], files={"file": ("s.txt", b"tiny", "text/plain")})
    assert r.status_code == 422
    # oversized
    r = client.post("/api/documents/upload", headers=tokens["admin_h"], files={"file": ("big.txt", b"a" * (16 * 1024 * 1024), "text/plain")})
    assert r.status_code == 413


def test_bulk_ingestion_is_idempotent_and_preserves_seed_corpus(tmp_path):
    """Re-scanning a directory must not duplicate documents, and deleting the
    DB row of an ingested file must not destroy the source file on disk."""
    from app.core.config import get_settings
    from app.core.database import db_session
    from app.models.document import Document
    from app.rag.ingestion.parsers import parse_bytes
    from app.services import document_service
    from sqlalchemy import select

    doc_text = (
        "# KB-900 Rotating Service Credentials\n\n"
        "## 1 When\nRotate credentials whenever an operator leaves the team or a secret leaks.\n\n"
        "## 2 How\nCreate the new secret, roll the deployment, then revoke the old secret after 24 hours.\n"
    )
    d = tmp_path / "docs"
    d.mkdir()
    (d / "KB-900.md").write_text(doc_text, encoding="utf-8")

    with db_session() as db:
        first = document_service.ingest_directory(db, d)
        assert first.stored == 1
        doc = db.scalar(select(Document).where(Document.name == "KB-900.md"))
        assert doc is not None and doc.content_hash

        second = document_service.ingest_directory(db, d)
        assert second.stored == 0
        assert any("already ingested" in s for s in second.skipped)
        assert db.scalars(select(Document).where(Document.name == "KB-900.md")).all().__len__() == 1

        # same content under a different file name is still a duplicate
        (d / "KB-900-renamed.md").write_text(doc_text, encoding="utf-8")
        third = document_service.ingest_directory(db, d)
        assert third.stored == 0

        document_service.delete_document(db, doc)
        assert db.scalar(select(Document).where(Document.id == doc.id)) is None
        assert (d / "KB-900.md").exists()
